"""
Live trade History — what each Live trade did and why.

- signal_breakdown(): entry / SL / TP derivation for one signal, resolved with the
  same per-pattern settings the strategy modules use (hammer classic vs inverted,
  doji single side, Hammer with candles BUY vs SELL rows, 35% pullback).
- HistoryStore: reads every logs/live/accounts/*/slots/*/live_trades.csv.
- format_trade_detail(): plain-text breakdown for the History window.

No Qt imports here so the maths stays unit-testable.
"""
from __future__ import annotations

import csv
import glob
import os
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple

import hammer_context_logic
import logic
from live_journal import TRADES_CSV_NAME, live_logs_dir


TRADE_EVENTS = ("SIGNAL", "DRY_RUN", "ORDER", "EXIT")
ERROR_EVENTS = ("ORDER_FAIL", "ERROR", "STARTUP_FAIL", "SAFETY_BLOCK")
SKIP_EVENTS = ("IGNORED", "SKIP_DUP", "SKIP_MAX_POS", "INDICATOR_BLOCK")

EVENT_HELP = {
    "SIGNAL": "Pattern found on the closed candle — levels computed from the preset.",
    "DRY_RUN": "Dry run is on — the order was NOT sent to MT5.",
    "ORDER": "Order accepted by MT5.",
    "EXIT": "Position closed in MT5 (SL, TP or manual).",
    "ORDER_FAIL": "MT5 rejected the order or no price was available.",
    "ERROR": "Live loop error (data / connection).",
    "STARTUP_FAIL": "Live could not start for this slot.",
    "SAFETY_BLOCK": "Blocked by a Live safety limit (daily trades / loss / spread / lot cap).",
    "IGNORED": "Pattern found but rejected by the strategy (e.g. risk above max SL).",
    "SKIP_DUP": "Already traded this signal candle.",
    "SKIP_MAX_POS": "Max open positions reached for this slot.",
    "INDICATOR_BLOCK": "Pattern found but an indicator filter blocked it.",
}

_RULE_SOURCE = {
    logic.EntryRule.NEXT_CANDLE_OPEN: "next candle open",
    logic.EntryRule.NEXT_CANDLE_CLOSE: "next candle close",
    logic.EntryRule.HAMMER_CLOSE: "signal candle close",
    logic.EntryRule.HAMMER_HIGH: "signal candle high",
    logic.EntryRule.HAMMER_LOW: "signal candle low",
}


@dataclass(frozen=True)
class SideParams:
    """Entry/SL settings that actually apply to one signal's side."""
    entry_rule: logic.EntryRule
    entry_offset: float
    sl_mode: logic.StopLossMode
    buffer_mode: logic.BufferMode
    sl_buffer_pct: float
    sl_buffer_flat: float
    sl_fixed_distance: float


def _f(value, default: float = 0.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def side_params(strategy_config, pattern_type: str, sig) -> Optional[SideParams]:
    """
    Mirror of each pattern's calculate_entry_price / calculate_stop_loss field choice.

    Hammer: inverted variant reads inverted_* (regardless of BUY/SELL).
    Doji: one row of settings for both directions.
    Hammer with candles (+35%): BUY uses the classic row, SELL the inverted row.
    """
    if strategy_config is None or sig is None:
        return None
    is_buy = sig.direction == logic.TradeDirection.BUY
    if hammer_context_logic.is_context_pattern_type(pattern_type):
        cfg = (
            strategy_config.to_buy_trade_config() if is_buy
            else strategy_config.to_sell_trade_config()
        )
        inverted = not is_buy
    elif pattern_type == "doji":
        cfg = strategy_config
        inverted = False
    else:
        cfg = strategy_config
        inverted = str(getattr(sig, "pattern_variant", "") or "").upper() == "INVERTED"

    if inverted:
        return SideParams(
            entry_rule=logic.coerce_entry_rule(getattr(cfg, "inverted_entry_rule", cfg.entry_rule)),
            entry_offset=_f(getattr(cfg, "inverted_entry_offset", cfg.entry_offset)),
            sl_mode=logic.coerce_stop_loss_mode(
                getattr(cfg, "inverted_sl_mode", logic.StopLossMode.CANDLE_EXTREME)
            ),
            buffer_mode=logic.coerce_buffer_mode(getattr(cfg, "inverted_buffer_mode", cfg.buffer_mode)),
            sl_buffer_pct=_f(getattr(cfg, "inverted_sl_buffer_pct", cfg.sl_buffer_pct)),
            sl_buffer_flat=_f(getattr(cfg, "inverted_sl_buffer_flat", cfg.sl_buffer_flat)),
            sl_fixed_distance=_f(getattr(cfg, "inverted_sl_fixed_distance", 0.0)),
        )
    return SideParams(
        entry_rule=logic.coerce_entry_rule(cfg.entry_rule),
        entry_offset=_f(getattr(cfg, "entry_offset", 0.0)),
        sl_mode=logic.coerce_stop_loss_mode(
            getattr(cfg, "sl_mode", logic.StopLossMode.CANDLE_EXTREME)
        ),
        buffer_mode=logic.coerce_buffer_mode(cfg.buffer_mode),
        sl_buffer_pct=_f(getattr(cfg, "sl_buffer_pct", 0.0)),
        sl_buffer_flat=_f(getattr(cfg, "sl_buffer_flat", 0.0)),
        sl_fixed_distance=_f(getattr(cfg, "sl_fixed_distance", 0.0)),
    )


@dataclass(frozen=True)
class _SideProbe:
    direction: Any
    pattern_variant: str


def configured_entry_offsets(strategy_config, pattern_type: str) -> List[Tuple[str, float]]:
    """Non-zero Entry Offset ($) rows of a preset, labelled by the row they come from."""
    if strategy_config is None:
        return []
    buy, sell = logic.TradeDirection.BUY, logic.TradeDirection.SELL
    if hammer_context_logic.is_context_pattern_type(pattern_type):
        probes = [("BUY row", buy, ""), ("SELL row", sell, "")]
    elif pattern_type == "doji":
        probes = [("Entry", buy, "")]
    else:
        probes = [("Classic", buy, "CLASSIC"), ("Inverted", buy, "INVERTED")]
    out: List[Tuple[str, float]] = []
    for label, direction, variant in probes:
        sig = _SideProbe(direction=direction, pattern_variant=variant)
        try:
            p = side_params(strategy_config, pattern_type, sig)
        except Exception:
            continue
        if p is not None and abs(p.entry_offset) > 1e-9:
            out.append((label, p.entry_offset))
    return out


_ORDER_MODE_TEXT = {
    "market": "Market",
    "limit_offset": "Limit — entry price ± offset (points)",
}


def entry_offset_ignored_warning(
    strategy_config,
    pattern_type: str,
    order_mode: str,
    limit_offset_from_market: bool = True,
) -> str:
    """
    Warning text when the preset has an Entry Offset ($) that the Live order type will not use.

    Market fills at ask/bid; 'Limit ± offset' from bid/ask ignores the strategy entry.
    'Limit — at strategy entry price' places the order at entry + offset (used).
    Hammer with candle 35% limit modes place the pullback limit from the strategy entry (used).
    Returns "" when nothing is ignored.
    """
    mode = (order_mode or "market").strip().lower()
    if mode not in _ORDER_MODE_TEXT:
        return ""
    if mode == "limit_offset":
        if not limit_offset_from_market or hammer_context_logic.is_35_pattern_type(pattern_type):
            return ""
    rows = configured_entry_offsets(strategy_config, pattern_type)
    if not rows:
        return ""
    offs = ", ".join(f"{label} ${value:+g}" for label, value in rows)
    how = (
        "Market orders fill at the current ask/bid"
        if mode == "market"
        else "the limit is placed from the current bid/ask"
    )
    return (
        f"WARNING: Entry Offset is set ({offs}) but Order type is "
        f"'{_ORDER_MODE_TEXT[mode]}' — {how}, so Live will NOT use the Entry Offset "
        f"and results can differ from the backtest. "
        f"Use Order type 'Limit — at strategy entry price' to apply it."
    )


def buffer_setting_text(p: SideParams) -> str:
    if p.sl_mode == logic.StopLossMode.FIXED_FROM_ENTRY:
        return f"fixed ${max(0.0, p.sl_fixed_distance):g} from entry"
    if p.buffer_mode == logic.BufferMode.FLAT_AMOUNT:
        return f"flat ${p.sl_buffer_flat:g}"
    if p.buffer_mode == logic.BufferMode.PERCENT_OF_RANGE:
        return f"{p.sl_buffer_pct:g}% of candle range"
    if p.buffer_mode == logic.BufferMode.PERCENT_OF_PRICE:
        return f"{p.sl_buffer_pct:g}% of price"
    return "none"


def _rule_price(rule: logic.EntryRule, signal_candle, entry_candle) -> Optional[float]:
    try:
        if rule == logic.EntryRule.NEXT_CANDLE_OPEN:
            return float(entry_candle.open)
        if rule == logic.EntryRule.NEXT_CANDLE_CLOSE:
            return float(entry_candle.close)
        if rule == logic.EntryRule.HAMMER_CLOSE:
            return float(signal_candle.close)
        if rule == logic.EntryRule.HAMMER_HIGH:
            return float(signal_candle.high)
        if rule == logic.EntryRule.HAMMER_LOW:
            return float(signal_candle.low)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def signal_breakdown(sig, strategy_config, pattern_type: str) -> Dict[str, Any]:
    """CSV fields explaining how entry / SL / TP came out of the preset for this signal."""
    out: Dict[str, Any] = {}
    if sig is None:
        return out
    c = getattr(sig, "hammer_candle", None)
    e = getattr(sig, "entry_candle", None)
    if c is not None and hasattr(c, "open"):
        out.update(
            signal_open=round(float(c.open), 5),
            signal_high=round(float(c.high), 5),
            signal_low=round(float(c.low), 5),
            signal_close=round(float(c.close), 5),
        )
    try:
        p = side_params(strategy_config, pattern_type, sig)
    except Exception:
        p = None
    if p is None or c is None or not hasattr(c, "low"):
        return out

    is_buy = sig.direction == logic.TradeDirection.BUY
    await_limit = bool(getattr(sig, "await_limit_fill", False))
    raw_signal_entry = getattr(sig, "signal_entry_price", None)
    signal_entry = float(raw_signal_entry) if await_limit and raw_signal_entry is not None else float(sig.entry_price)
    sl = float(sig.stop_loss)

    rule_px = _rule_price(p.entry_rule, c, e)
    out.update(
        entry_rule=p.entry_rule.value,
        entry_rule_price=round(rule_px, 5) if rule_px is not None else "",
        entry_offset=p.entry_offset,
        signal_entry=round(signal_entry, 5),
        pullback_pct=_f(getattr(sig, "entry_pullback_pct", 0.0)) if await_limit else "",
        sl_mode=p.sl_mode.value,
        sl_buffer_mode=(
            p.sl_mode.value if p.sl_mode == logic.StopLossMode.FIXED_FROM_ENTRY else p.buffer_mode.value
        ),
        sl_buffer_setting=buffer_setting_text(p),
    )
    if p.sl_mode == logic.StopLossMode.FIXED_FROM_ENTRY:
        anchor, label = signal_entry, "signal entry"
    else:
        anchor = float(c.low) if is_buy else float(c.high)
        label = "candle low" if is_buy else "candle high"
    buffer_amount = (anchor - sl) if is_buy else (sl - anchor)
    out.update(
        sl_anchor_label=label,
        sl_anchor=round(anchor, 5),
        sl_buffer_amount=round(buffer_amount, 5),
    )
    return out


# ---------------------------------------------------------------------------
# Reading journals
# ---------------------------------------------------------------------------

def _label_from_dir(dirname: str) -> str:
    """'Main_ab12cd34' → 'Main' (folder names are <label>_<id>)."""
    base = os.path.basename(dirname.rstrip(os.sep))
    head, sep, _tail = base.rpartition("_")
    return head if sep and head else base


class HistoryStore:
    """Cached reader over all slot journals (re-reads a file only when it changes)."""

    def __init__(self, app_root: str):
        self.app_root = app_root
        self._cache: Dict[str, Tuple[float, int, List[Dict[str, str]]]] = {}

    def journal_paths(self) -> List[str]:
        """Desk slots (accounts/*/slots/*) and the API desk (api/...)."""
        root = live_logs_dir(self.app_root)
        return sorted(glob.glob(os.path.join(root, "**", TRADES_CSV_NAME), recursive=True))

    def _read(self, path: str) -> List[Dict[str, str]]:
        try:
            st = os.stat(path)
        except OSError:
            return []
        hit = self._cache.get(path)
        if hit is not None and hit[0] == st.st_mtime and hit[1] == st.st_size:
            return hit[2]
        slot_dir = os.path.dirname(path)
        if os.path.basename(os.path.dirname(slot_dir)) == "slots":
            account_label = _label_from_dir(os.path.dirname(os.path.dirname(slot_dir)))
            slot_label = _label_from_dir(slot_dir)
        else:
            account_label = os.path.basename(slot_dir) or "Live"
            slot_label = ""
        rows: List[Dict[str, str]] = []
        try:
            with open(path, "r", encoding="utf-8", newline="") as f:
                for row in csv.DictReader(f):
                    row = {k: (v if v is not None else "") for k, v in row.items() if k is not None}
                    if not row.get("account"):
                        row["account"] = account_label
                    if not row.get("slot"):
                        row["slot"] = slot_label
                    row["_journal"] = path
                    rows.append(row)
        except (OSError, csv.Error, UnicodeDecodeError):
            return hit[2] if hit else []
        self._cache[path] = (st.st_mtime, st.st_size, rows)
        return rows

    def load(self) -> List[Dict[str, str]]:
        """All rows across accounts/slots, newest first."""
        rows: List[Dict[str, str]] = []
        paths = self.journal_paths()
        for stale in set(self._cache) - set(paths):
            self._cache.pop(stale, None)
        for path in paths:
            rows.extend(self._read(path))
        # Same-second rows keep file order (SIGNAL before ORDER), then flip to newest first.
        order = sorted(range(len(rows)), key=lambda i: (rows[i].get("logged_at") or "", i), reverse=True)
        return [rows[i] for i in order]


def event_category(event: str) -> str:
    ev = (event or "").upper()
    if ev in ERROR_EVENTS:
        return "error"
    if ev in SKIP_EVENTS:
        return "skip"
    return "trade"


def trade_group_key(row: Dict[str, str]) -> Tuple[str, str, str]:
    """Rows of the same signal (SIGNAL → ORDER → EXIT) share journal + signal bar + side."""
    return (
        row.get("_journal", ""),
        row.get("signal_bar_time", ""),
        row.get("direction", ""),
    )


def related_rows(rows: List[Dict[str, str]], row: Dict[str, str]) -> List[Dict[str, str]]:
    key = trade_group_key(row)
    if not key[1]:
        return [row]
    same = [r for r in rows if trade_group_key(r) == key]
    return sorted(same, key=lambda r: r.get("logged_at") or "")


# ---------------------------------------------------------------------------
# Detail text
# ---------------------------------------------------------------------------

def _num(row: Dict[str, str], key: str) -> Optional[float]:
    raw = row.get(key, "")
    if raw in ("", None):
        return None
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


def _px(v: Optional[float]) -> str:
    return "—" if v is None else f"{v:.2f}"


def _is_buy(row: Dict[str, str]) -> bool:
    return (row.get("direction") or "").upper() == "BUY"


def entry_calc_text(row: Dict[str, str]) -> str:
    """'next candle open 2652.00 + 0.30 = 2652.30' (+ pullback step for 35%)."""
    rule = row.get("entry_rule", "")
    rule_px = _num(row, "entry_rule_price")
    offset = _num(row, "entry_offset")
    signal_entry = _num(row, "signal_entry")
    if not rule or rule_px is None:
        return ""
    try:
        src = _RULE_SOURCE.get(logic.coerce_entry_rule(rule), rule)
    except Exception:
        src = rule
    text = f"{src} {_px(rule_px)}"
    if offset is not None:
        text += f" {'+' if offset >= 0 else '−'} {abs(offset):g}"
    text += f" = {_px(signal_entry)}"
    pull = _num(row, "pullback_pct")
    strat_entry = _num(row, "strategy_entry")
    if pull is not None and strat_entry is not None:
        text += f" → pullback {pull:g}% → limit {_px(strat_entry)}"
    return text


def sl_calc_text(row: Dict[str, str]) -> str:
    """'candle low 2640.00 − 0.50 (flat $0.5) = 2639.50'."""
    anchor = _num(row, "sl_anchor")
    buf = _num(row, "sl_buffer_amount")
    sl = _num(row, "strategy_sl")
    if anchor is None or buf is None or sl is None:
        return ""
    label = row.get("sl_anchor_label", "") or ("candle low" if _is_buy(row) else "candle high")
    op = "−" if _is_buy(row) else "+"
    setting = row.get("sl_buffer_setting", "")
    return f"{label} {_px(anchor)} {op} {buf:.2f}" + (f" ({setting})" if setting else "") + f" = {_px(sl)}"


def tp_calc_text(row: Dict[str, str]) -> str:
    """'risk 12.80 × RR 2 = 25.60 → TP 2677.90' (from signal entry)."""
    rr = _num(row, "rr_multiple")
    sl = _num(row, "strategy_sl")
    tp = _num(row, "strategy_tp")
    base = _num(row, "signal_entry")
    if base is None:
        base = _num(row, "strategy_entry")
    if rr is None or sl is None or tp is None or base is None:
        return ""
    risk = abs(base - sl)
    return f"risk {risk:.2f} × RR {rr:g} = {risk * rr:.2f} → TP {_px(tp)}"


def sl_check_text(row: Dict[str, str]) -> str:
    """OK when the SL sent to MT5 equals the strategy SL; MOVED otherwise."""
    if (row.get("event") or "").upper() not in ("ORDER", "EXIT"):
        return ""
    sent = _num(row, "stop_loss")
    strat = _num(row, "strategy_sl")
    if sent is None or strat is None:
        return ""
    return "OK" if abs(sent - strat) < 1e-6 else f"MOVED {sent - strat:+.2f}"


def format_trade_detail(row: Dict[str, str], timeline: Optional[List[Dict[str, str]]] = None) -> str:
    ev = (row.get("event") or "").upper()
    direction = (row.get("direction") or "").upper()
    is_buy = direction == "BUY"
    lines: List[str] = []

    head = f"{ev}"
    if direction:
        head += f" — {direction}"
    head += f"  {row.get('symbol', '')} {row.get('timeframe', '')}"
    pattern = row.get("pattern", "")
    variant = row.get("pattern_variant", "")
    if pattern:
        head += f"  |  {pattern}" + (f" ({variant})" if variant else "")
    lines.append(head)
    if ev in EVENT_HELP:
        lines.append(f"  {EVENT_HELP[ev]}")
    lines.append(
        f"Account: {row.get('account', '') or '—'}   Slot: {row.get('slot', '') or '—'}   "
        f"Preset: {row.get('preset', '') or '—'}   Magic: {row.get('magic', '') or '—'}"
    )
    lines.append(
        f"Logged: {row.get('logged_at', '')}   Dry run: {row.get('dry_run', '') or '—'}   "
        f"Order type: {row.get('order_mode', '') or '—'}   Lots: {row.get('volume', '') or '—'}"
    )
    if row.get("signal_bar_time"):
        lines.append(
            f"Signal candle: {row.get('signal_bar_time')}   Entry candle: {row.get('entry_bar_time', '') or '—'}"
        )

    why = (row.get("reason") or "").strip()
    mt5_msg = (row.get("mt5_message") or "").strip()
    if event_category(ev) != "trade" or (why and ev not in ("SIGNAL",)):
        lines.append("")
        title = "ERROR / BLOCK" if event_category(ev) == "error" else "REASON"
        lines.append(title)
        if why:
            lines.append(f"  {why}")
        if mt5_msg and mt5_msg != why:
            lines.append(f"  MT5: {mt5_msg}")

    o, h, lo, c = (_num(row, k) for k in ("signal_open", "signal_high", "signal_low", "signal_close"))
    if None not in (o, h, lo, c):
        lines.append("")
        lines.append(f"SIGNAL CANDLE   O {_px(o)}   H {_px(h)}   L {_px(lo)}   C {_px(c)}")

    strat_entry = _num(row, "strategy_entry")
    strat_sl = _num(row, "strategy_sl")
    strat_tp = _num(row, "strategy_tp")
    sent_entry = _num(row, "entry_price")
    sent_sl = _num(row, "stop_loss")
    sent_tp = _num(row, "target")
    signal_entry = _num(row, "signal_entry")
    rule_px = _num(row, "entry_rule_price")
    offset = _num(row, "entry_offset")
    pull = _num(row, "pullback_pct")

    if strat_entry is not None or rule_px is not None:
        lines.append("")
        lines.append("ENTRY")
        rule = row.get("entry_rule", "")
        if rule:
            try:
                src = _RULE_SOURCE.get(logic.coerce_entry_rule(rule), rule)
            except Exception:
                src = rule
            lines.append(f"  Rule {rule}: {src} = {_px(rule_px)}")
        if offset is not None and rule_px is not None:
            sign = "+" if offset >= 0 else "−"
            lines.append(
                f"  {sign} entry offset {abs(offset):g} = {_px(signal_entry)}  (strategy entry)"
            )
        if pull is not None and strat_entry is not None:
            toward = "down" if is_buy else "up"
            lines.append(
                f"  Pullback {pull:g}% of (entry − SL) toward SL ({toward}) → limit {_px(strat_entry)}"
            )
        if sent_entry is not None and ev in ("ORDER", "EXIT") and strat_entry is not None:
            delta = sent_entry - strat_entry
            lines.append(f"  Sent to MT5: {_px(sent_entry)}  (Δ {delta:+.2f} vs strategy entry — market fill)")

    if strat_sl is not None:
        lines.append("")
        lines.append("STOP LOSS")
        anchor = _num(row, "sl_anchor")
        buf = _num(row, "sl_buffer_amount")
        label = row.get("sl_anchor_label", "") or ("candle low" if is_buy else "candle high")
        setting = row.get("sl_buffer_setting", "")
        if anchor is not None and buf is not None:
            op = "−" if is_buy else "+"
            lines.append(
                f"  {label} {_px(anchor)} {op} buffer {buf:.2f}"
                + (f" ({setting})" if setting else "")
                + f" = SL {_px(strat_sl)}"
            )
            if buf < -1e-9:
                lines.append("  ⚠ buffer is negative — SL is INSIDE the candle (check preset).")
        else:
            lines.append(f"  SL {_px(strat_sl)}")
        if sent_sl is not None and ev in ("ORDER", "EXIT"):
            ok = abs(sent_sl - strat_sl) < 1e-6
            lines.append(
                f"  Sent to MT5: {_px(sent_sl)}  "
                + ("✓ same as strategy SL (locked, not slid with fill)" if ok
                   else f"✗ differs from strategy SL by {sent_sl - strat_sl:+.2f}")
            )

    if strat_tp is not None:
        lines.append("")
        lines.append("TARGET")
        rr = _num(row, "rr_multiple")
        base = signal_entry if signal_entry is not None else strat_entry
        if rr is not None and base is not None and strat_sl is not None:
            risk = abs(base - strat_sl)
            op = "+" if is_buy else "−"
            lines.append(
                f"  risk {risk:.2f} (entry {_px(base)} ↔ SL {_px(strat_sl)}) × RR {rr:g} = {risk * rr:.2f}"
            )
            lines.append(f"  TP = {_px(base)} {op} {risk * rr:.2f} = {_px(strat_tp)}")
        else:
            lines.append(f"  TP {_px(strat_tp)}")
        if sent_tp is not None and ev in ("ORDER", "EXIT"):
            if abs(sent_tp - strat_tp) < 1e-6:
                lines.append(f"  Sent to MT5: {_px(sent_tp)}  ✓ same as strategy TP")
            elif pull is not None:
                lines.append(f"  Sent to MT5: {_px(sent_tp)}  (pullback keeps TP from signal entry)")
            else:
                lines.append(
                    f"  Sent to MT5: {_px(sent_tp)}  (rebuilt from fill {_px(sent_entry)} × RR so reward stays RR × risk)"
                )

    exit_px = _num(row, "exit_price")
    profit = _num(row, "profit")
    if ev == "EXIT" or exit_px is not None or profit is not None:
        lines.append("")
        lines.append("RESULT")
        lines.append(
            f"  Exit {_px(exit_px)}   Profit {'—' if profit is None else f'{profit:+.2f}'}   "
            f"{row.get('close_reason', '')}"
        )

    params = (row.get("params") or "").strip()
    if params:
        lines.append("")
        lines.append("PRESET PARAMETERS USED")
        for part in params.split(" | "):
            part = part.strip()
            if part and part != "[PARAMS]":
                lines.append(f"  • {part}")

    if timeline and len(timeline) > 1:
        lines.append("")
        lines.append("TIMELINE (same signal candle)")
        for r in timeline:
            marker = "▶" if r is row else " "
            bits = [r.get("logged_at", ""), (r.get("event") or "").upper()]
            why_r = (r.get("reason") or r.get("mt5_message") or "").strip()
            if why_r:
                bits.append(why_r)
            lines.append(f" {marker} " + "  ".join(b for b in bits if b))

    return "\n".join(lines)


# ---------------------------------------------------------------------------
# One line per trade + Excel export
# ---------------------------------------------------------------------------

_PRIMARY_ORDER = (
    "ORDER", "DRY_RUN", "ORDER_FAIL", "SAFETY_BLOCK", "INDICATOR_BLOCK",
    "IGNORED", "SKIP_MAX_POS", "SKIP_DUP", "SIGNAL", "EXIT",
)

_STATUS = {
    "ORDER": "Open (sent to MT5)",
    "DRY_RUN": "Dry run (not sent)",
    "ORDER_FAIL": "Order failed",
    "SAFETY_BLOCK": "Blocked (safety limit)",
    "INDICATOR_BLOCK": "Blocked (indicator)",
    "IGNORED": "Rejected by strategy",
    "SKIP_MAX_POS": "Skipped (max positions)",
    "SKIP_DUP": "Skipped (already traded)",
    "SIGNAL": "Signal only",
    "EXIT": "Closed",
}


def trade_problem_text(rows: List[Dict[str, str]]) -> str:
    """First thing worth a look for one trade: error/block reason, moved SL, SL inside candle."""
    problems: List[str] = []
    for r in rows:
        ev = (r.get("event") or "").upper()
        if event_category(ev) == "error":
            problems.append(f"{ev}: {(r.get('reason') or r.get('mt5_message') or '').strip()}")
        check = sl_check_text(r)
        if check.startswith("MOVED"):
            problems.append(f"SL {check} vs strategy")
        buf = _num(r, "sl_buffer_amount")
        if buf is not None and buf < -1e-9:
            problems.append("SL inside candle (negative buffer)")
    seen: List[str] = []
    for p in problems:
        if p not in seen:
            seen.append(p)
    return "; ".join(seen)


def trade_summaries(rows: List[Dict[str, str]]) -> List[Dict[str, Any]]:
    """Collapse SIGNAL → ORDER → EXIT rows into one line per trade (newest first)."""
    groups: Dict[Tuple[str, str, str], List[Dict[str, str]]] = {}
    order: List[Tuple[str, str, str]] = []
    for r in rows:
        key = trade_group_key(r)
        ev = (r.get("event") or "").upper()
        if not key[1]:
            if ev != "EXIT":
                continue
            key = (key[0], f"exit:{r.get('logged_at', '')}:{r.get('mt5_order_id', '')}", key[2])
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(r)

    out: List[Dict[str, Any]] = []
    for key in order:
        grp = groups[key]
        by_event = {(r.get("event") or "").upper(): r for r in grp}
        primary = next((by_event[e] for e in _PRIMARY_ORDER if e in by_event), grp[0])
        exit_row = by_event.get("EXIT")
        order_row = by_event.get("ORDER")
        status = _STATUS["EXIT"] if exit_row else _STATUS.get(
            (primary.get("event") or "").upper(), primary.get("event", "")
        )
        levels = order_row or primary
        out.append({
            "Signal candle": primary.get("signal_bar_time", ""),
            "Logged": primary.get("logged_at", ""),
            "Account": primary.get("account", ""),
            "Slot": primary.get("slot", ""),
            "Preset": primary.get("preset", ""),
            "Pattern": primary.get("pattern", ""),
            "Variant": primary.get("pattern_variant", ""),
            "Symbol": primary.get("symbol", ""),
            "TF": primary.get("timeframe", ""),
            "Side": primary.get("direction", ""),
            "Status": status,
            "Lots": _num(levels, "volume"),
            "Candle high": _num(primary, "signal_high"),
            "Candle low": _num(primary, "signal_low"),
            "Entry calc": entry_calc_text(primary),
            "Strategy entry": _num(primary, "strategy_entry"),
            "Entry sent": _num(order_row, "entry_price") if order_row else None,
            "SL calc": sl_calc_text(primary),
            "SL buffer": _num(primary, "sl_buffer_amount"),
            "SL": _num(primary, "strategy_sl"),
            "SL sent": _num(order_row, "stop_loss") if order_row else None,
            "SL check": sl_check_text(order_row) if order_row else "",
            "TP calc": tp_calc_text(primary),
            "TP": _num(primary, "strategy_tp"),
            "TP sent": _num(order_row, "target") if order_row else None,
            "RR": _num(primary, "rr_multiple"),
            "Exit price": _num(exit_row, "exit_price") if exit_row else None,
            "Profit": _num(exit_row, "profit") if exit_row else None,
            "Close reason": (exit_row or {}).get("close_reason", ""),
            "Ticket": (order_row or exit_row or {}).get("mt5_order_id", ""),
            "Problem": trade_problem_text(grp),
            "Magic": primary.get("magic", ""),
        })
    return out


EVENT_EXPORT_HEADERS = {
    "logged_at": "Logged", "event": "Event", "symbol": "Symbol", "timeframe": "TF",
    "pattern": "Pattern", "pattern_variant": "Variant", "direction": "Side", "volume": "Lots",
    "entry_price": "Entry sent", "stop_loss": "SL sent", "target": "TP sent", "risk": "Risk",
    "rr_multiple": "RR", "dry_run": "Dry run", "order_mode": "Order type", "magic": "Magic",
    "mt5_order_id": "Ticket", "mt5_message": "MT5 message", "reason": "Reason / error",
    "signal_bar_time": "Signal candle", "entry_bar_time": "Entry candle",
    "exit_price": "Exit price", "profit": "Profit", "close_reason": "Close reason",
    "slot": "Slot", "strategy_entry": "Strategy entry", "strategy_sl": "Strategy SL",
    "strategy_tp": "Strategy TP", "account": "Account", "preset": "Preset",
    "signal_open": "Candle open", "signal_high": "Candle high", "signal_low": "Candle low",
    "signal_close": "Candle close", "entry_rule": "Entry rule",
    "entry_rule_price": "Entry rule price", "entry_offset": "Entry offset",
    "signal_entry": "Signal entry", "pullback_pct": "Pullback %", "sl_mode": "SL mode",
    "sl_anchor_label": "SL from", "sl_anchor": "SL anchor price",
    "sl_buffer_mode": "Buffer mode", "sl_buffer_setting": "Buffer setting",
    "sl_buffer_amount": "SL buffer", "params": "Preset parameters",
}

_NUMERIC_KEYS = {
    "volume", "entry_price", "stop_loss", "target", "risk", "rr_multiple", "exit_price",
    "profit", "strategy_entry", "strategy_sl", "strategy_tp", "signal_open", "signal_high",
    "signal_low", "signal_close", "entry_rule_price", "entry_offset", "signal_entry",
    "pullback_pct", "sl_anchor", "sl_buffer_amount",
}


def _excel_value(key: str, raw: Any) -> Any:
    if key in _NUMERIC_KEYS:
        try:
            return float(raw) if raw not in ("", None) else None
        except (TypeError, ValueError):
            return raw
    return raw


def export_history_xlsx(path: str, rows: List[Dict[str, str]], *, filters_text: str = "") -> Dict[str, int]:
    """
    Shareable workbook: README, Trades (one line per trade with entry/SL/TP maths),
    Errors (errors + safety blocks), All_Events (every journal column).
    Returns sheet → row count.
    """
    import openpyxl
    from datetime import datetime
    from openpyxl.styles import Alignment, Font, PatternFill
    from live_journal import TRADE_CSV_COLUMNS

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="34A853")
    problem_fill = PatternFill("solid", fgColor="FDECEA")
    win_font = Font(color="137333", bold=True)
    loss_font = Font(color="C5221F", bold=True)

    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "Hammer Live — Trade History"
    readme["A1"].font = Font(bold=True, size=14)
    info = [
        ("Exported", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Filters", filters_text or "none (all events)"),
        ("Trades", "One line per signal: status, entry / SL / TP maths, what was sent to MT5, exit and profit. "
                   "'Problem' shows errors, blocks or a moved SL — start there when checking."),
        ("SL calc", "BUY: candle low − buffer. SELL: candle high + buffer. 'SL check' = OK when MT5 got the same SL."),
        ("Entry calc", "Price picked by the entry rule ± entry offset; 35% pattern then pulls back toward SL."),
        ("Errors", "Order failures, Live errors, startup failures and safety blocks with the MT5 message."),
        ("All_Events", "Every Live journal row (signals, dry runs, orders, exits, skips) with all columns."),
    ]
    for i, (k, v) in enumerate(info, start=3):
        readme.cell(row=i, column=1, value=k).font = Font(bold=True)
        readme.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 14
    readme.column_dimensions["B"].width = 110

    def write_sheet(name: str, headers: List[str], data: List[List[Any]], *,
                    highlight: Optional[List[bool]] = None, profit_col: Optional[int] = None) -> int:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for i, values in enumerate(data):
            ws.append(values)
            r = ws.max_row
            if highlight is not None and highlight[i]:
                for cell in ws[r]:
                    cell.fill = problem_fill
            if profit_col is not None and isinstance(values[profit_col], (int, float)):
                ws.cell(row=r, column=profit_col + 1).font = win_font if values[profit_col] >= 0 else loss_font
        ws.freeze_panes = "A2"
        if data:
            ws.auto_filter.ref = ws.dimensions
        for col_cells in ws.columns:
            width = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(60, max(8, width + 2))
        return len(data)

    counts: Dict[str, int] = {}
    trades = trade_summaries(rows)
    t_headers = list(trades[0].keys()) if trades else ["(no trades)"]
    counts["Trades"] = write_sheet(
        "Trades", t_headers, [[t[h] for h in t_headers] for t in trades],
        highlight=[bool(t["Problem"]) for t in trades],
        profit_col=t_headers.index("Profit") if trades else None,
    )

    err_keys = [
        "logged_at", "account", "slot", "preset", "pattern", "timeframe", "event", "direction",
        "reason", "mt5_message", "signal_bar_time", "strategy_entry", "strategy_sl", "strategy_tp",
    ]
    errors = [r for r in rows if event_category(r.get("event", "")) == "error"]
    counts["Errors"] = write_sheet(
        "Errors", [EVENT_EXPORT_HEADERS.get(k, k) for k in err_keys],
        [[_excel_value(k, r.get(k, "")) for k in err_keys] for r in errors],
    )

    ev_keys = list(TRADE_CSV_COLUMNS)
    counts["All_Events"] = write_sheet(
        "All_Events", [EVENT_EXPORT_HEADERS.get(k, k) for k in ev_keys],
        [[_excel_value(k, r.get(k, "")) for k in ev_keys] for r in rows],
        highlight=[event_category(r.get("event", "")) == "error" for r in rows],
        profit_col=ev_keys.index("profit"),
    )
    wb.save(path)
    return counts
