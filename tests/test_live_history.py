"""
History (menu bar → History): per-trade entry / SL + buffer / TP breakdown for every
pattern, journal header upgrade, multi-slot reader, and detail text.
"""

from __future__ import annotations

import csv
import os
from types import SimpleNamespace
from unittest.mock import MagicMock

import doji_logic
import hammer_context_logic as hc
import live_history
import live_journal as lj
import logic
from live import LiveTradingEngine
from strategy import hammer_with_candles_35_logic as hwc35
from strategy import hammer_with_candles_logic as hwc


def _c(ts: str, o: float, h: float, l: float, c: float) -> logic.Candle:
    return logic.Candle(ts, o, h, l, c)


def _sig(direction, signal, nxt, entry, sl, *, variant=None, rr=2.0, **kw):
    risk = abs(entry - sl)
    tp = entry + rr * risk if direction == logic.TradeDirection.BUY else entry - rr * risk
    return logic.TradeSignal(
        direction=direction, hammer_candle=signal, entry_candle=nxt,
        entry_price=entry, stop_loss=sl, risk=risk, rr_multiple=rr, target=tp,
        timeframe="3m", pattern_variant=variant, **kw,
    )


# ---------------------------------------------------------------------------
# side_params: the settings row each pattern really uses
# ---------------------------------------------------------------------------

def test_hammer_inverted_buy_uses_inverted_row_not_direction():
    cfg = logic.StrategyConfig(
        entry_offset=0.10, sl_buffer_flat=0.20, buffer_mode=logic.BufferMode.FLAT_AMOUNT,
        inverted_entry_offset=0.70, inverted_sl_buffer_flat=0.90,
        inverted_buffer_mode=logic.BufferMode.FLAT_AMOUNT,
    )
    s = _c("s", 100, 104, 99.5, 100.2)
    n = _c("n", 100.3, 101, 100, 100.8)
    sig = _sig(logic.TradeDirection.BUY, s, n, 101.0, 98.6, variant="INVERTED")
    p = live_history.side_params(cfg, "hammer", sig)
    assert p.entry_offset == 0.70 and p.sl_buffer_flat == 0.90
    classic = _sig(logic.TradeDirection.BUY, s, n, 101.0, 98.6, variant="CLASSIC")
    p2 = live_history.side_params(cfg, "hammer", classic)
    assert p2.entry_offset == 0.10 and p2.sl_buffer_flat == 0.20


def test_hwc_sell_uses_inverted_row_buy_uses_classic_row():
    cfg = hwc.prepare_config(hwc.HammerContextConfig(
        entry_offset=0.30, sl_buffer_flat=0.50, buffer_mode=logic.BufferMode.FLAT_AMOUNT,
        inverted_entry_offset=-0.40, inverted_sl_buffer_flat=1.25,
        inverted_buffer_mode=logic.BufferMode.FLAT_AMOUNT,
    ))
    s = _c("s", 100, 101, 90, 99)
    n = _c("n", 99, 100, 98, 99.5)
    sell = _sig(logic.TradeDirection.SELL, s, n, 98.6, 102.25, variant="INVERTED")
    buy = _sig(logic.TradeDirection.BUY, s, n, 99.3, 89.5, variant="CLASSIC")
    assert live_history.side_params(cfg, hc.PATTERN_TYPE, sell).sl_buffer_flat == 1.25
    assert live_history.side_params(cfg, hc.PATTERN_TYPE, sell).entry_offset == -0.40
    assert live_history.side_params(cfg, hc.PATTERN_TYPE, buy).sl_buffer_flat == 0.50


def test_doji_single_row_both_directions():
    cfg = doji_logic.DojiStrategyConfig(
        entry_offset=0.25, buffer_mode=logic.BufferMode.FLAT_AMOUNT, sl_buffer_flat=0.40,
    )
    s = _c("s", 100.0, 101.0, 99.0, 100.05)
    n = _c("n", 100.2, 101.0, 100.0, 100.5)
    sell = _sig(logic.TradeDirection.SELL, s, n, 100.45, 101.4, variant="INVERTED")
    p = live_history.side_params(cfg, "doji", sell)
    assert p.entry_offset == 0.25 and p.sl_buffer_flat == 0.40


# ---------------------------------------------------------------------------
# signal_breakdown numbers — built from the real strategy functions
# ---------------------------------------------------------------------------

def test_hammer_buy_breakdown_low_minus_flat_buffer():
    hammer = _c("s", 2650, 2660, 2640, 2655)
    nxt = _c("e", 2652, 2665, 2650, 2660)
    cfg = logic.StrategyConfig(
        entry_rule=logic.EntryRule.NEXT_CANDLE_OPEN, entry_offset=0.30,
        buffer_mode=logic.BufferMode.FLAT_AMOUNT, sl_buffer_flat=0.50,
        enable_risk_limit=False,
    )
    entry = logic.calculate_entry_price(hammer, nxt, cfg)
    sl = logic.calculate_stop_loss(hammer, logic.TradeDirection.BUY, cfg, entry_price=entry)
    sig = _sig(logic.TradeDirection.BUY, hammer, nxt, entry, sl, variant="CLASSIC")
    b = live_history.signal_breakdown(sig, cfg, "hammer")
    assert b["entry_rule_price"] == 2652.0
    assert b["entry_offset"] == 0.30
    assert abs(b["signal_entry"] - 2652.30) < 1e-9
    assert b["sl_anchor_label"] == "candle low" and b["sl_anchor"] == 2640.0
    assert abs(b["sl_buffer_amount"] - 0.50) < 1e-9
    assert b["sl_buffer_setting"] == "flat $0.5"


def test_hammer_sell_percent_of_range_buffer_amount():
    hammer = _c("s", 2655, 2660, 2640, 2650)
    nxt = _c("e", 2652, 2658, 2648, 2650)
    cfg = logic.StrategyConfig(
        inverted_entry_rule=logic.EntryRule.NEXT_CANDLE_OPEN,
        inverted_buffer_mode=logic.BufferMode.PERCENT_OF_RANGE,
        inverted_sl_buffer_pct=10.0, enable_risk_limit=False,
    )
    v = logic.HammerVariant.INVERTED
    entry = logic.calculate_entry_price(hammer, nxt, cfg, v)
    sl = logic.calculate_stop_loss(hammer, logic.TradeDirection.SELL, cfg, v, entry_price=entry)
    sig = _sig(logic.TradeDirection.SELL, hammer, nxt, entry, sl, variant="INVERTED")
    b = live_history.signal_breakdown(sig, cfg, "hammer")
    assert b["sl_anchor_label"] == "candle high" and b["sl_anchor"] == 2660.0
    assert abs(b["sl_buffer_amount"] - 2.0) < 1e-9  # 10% of 20 range
    assert abs(sl - 2662.0) < 1e-9


def test_hwc35_breakdown_shows_signal_entry_and_pullback():
    signal = _c("s", 100.0, 103.5, 90.0, 102.0)
    nxt = _c("n", 102.0, 103.0, 101.0, 102.5)
    cfg = hwc35.prepare_config(hwc35.HammerContextConfig(
        lookback_candles=1, enable_sell=False, enable_risk_limit=False,
        entry_rule=logic.EntryRule.NEXT_CANDLE_OPEN, entry_offset=0.0,
        entry_pullback_pct=35.0, buffer_mode=logic.BufferMode.FLAT_AMOUNT, sl_buffer_flat=0.5,
    ))
    pads = [_c("p", 95.0, 96.0, 94.0, 95.0), signal, nxt]
    sig = [s for s in hwc35.run_strategy(pads, "3m", cfg) if not s.ignored][0]
    b = live_history.signal_breakdown(sig, cfg, hc.PATTERN_TYPE_35)
    assert b["signal_entry"] == 102.0
    assert b["pullback_pct"] == 35.0
    assert b["sl_anchor"] == 90.0 and abs(b["sl_buffer_amount"] - 0.5) < 1e-9
    # limit = 102 − 35% × (102 − 89.5)
    assert abs(sig.entry_price - (102.0 - 0.35 * 12.5)) < 1e-9


def test_fixed_from_entry_anchor_is_signal_entry():
    hammer = _c("s", 2650, 2660, 2640, 2655)
    nxt = _c("e", 2652, 2665, 2650, 2660)
    cfg = logic.StrategyConfig(
        sl_mode=logic.StopLossMode.FIXED_FROM_ENTRY, sl_fixed_distance=5.0,
        enable_risk_limit=False,
    )
    entry = logic.calculate_entry_price(hammer, nxt, cfg)
    sl = logic.calculate_stop_loss(hammer, logic.TradeDirection.BUY, cfg, entry_price=entry)
    sig = _sig(logic.TradeDirection.BUY, hammer, nxt, entry, sl, variant="CLASSIC")
    b = live_history.signal_breakdown(sig, cfg, "hammer")
    assert b["sl_anchor_label"] == "signal entry"
    assert abs(b["sl_buffer_amount"] - 5.0) < 1e-9
    assert b["sl_buffer_setting"] == "fixed $5 from entry"


# ---------------------------------------------------------------------------
# Journal + reader
# ---------------------------------------------------------------------------

def test_old_csv_header_upgraded_without_losing_rows(tmp_path):
    slot = str(tmp_path)
    old_cols = lj.TRADE_CSV_COLUMNS[: lj.TRADE_CSV_COLUMNS.index("strategy_tp") + 1]
    with open(lj.trades_csv_path(slot), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=old_cols)
        w.writeheader()
        w.writerow({"event": "ORDER", "symbol": "XAUUSD", "stop_loss": "2639.5"})
    lj.append_trade_row(slot, {"event": "EXIT", "symbol": "XAUUSD", "sl_buffer_amount": 0.5})
    with open(lj.trades_csv_path(slot), encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    assert [r["event"] for r in rows] == ["ORDER", "EXIT"]
    assert rows[0]["stop_loss"] == "2639.5"
    assert rows[1]["sl_buffer_amount"] == "0.5"


def test_store_reads_desk_slots_and_api_journal(tmp_path):
    root = str(tmp_path)
    s1 = lj.slot_journal_dir(root, "ab12", "Main", "s1", "Gold", "3m")
    lj.append_trade_row(s1, {"event": "ORDER", "logged_at": "2026-09-25T10:00:00"})
    api = os.path.join(lj.live_logs_dir(root), "api")
    lj.append_trade_row(api, {"event": "ERROR", "logged_at": "2026-09-25T11:00:00", "account": "Api1"})
    rows = live_history.HistoryStore(root).load()
    assert [r["event"] for r in rows] == ["ERROR", "ORDER"]  # newest first
    assert rows[0]["account"] == "Api1"
    assert rows[1]["account"] == "Main" and rows[1]["slot"] == "Gold_3m"


def test_detail_text_explains_sl_buffer_and_lock():
    row = {
        "event": "ORDER", "direction": "BUY", "symbol": "XAUUSD", "timeframe": "3m",
        "pattern": "Hammer", "account": "Main", "slot": "Gold", "preset": "gold.json",
        "signal_open": "2650", "signal_high": "2660", "signal_low": "2640", "signal_close": "2655",
        "entry_rule": "NEXT_CANDLE_OPEN", "entry_rule_price": "2652", "entry_offset": "0.3",
        "signal_entry": "2652.3", "strategy_entry": "2652.3", "strategy_sl": "2639.5",
        "strategy_tp": "2677.9", "entry_price": "2652.1", "stop_loss": "2639.5",
        "target": "2677.3", "rr_multiple": "2", "sl_anchor_label": "candle low",
        "sl_anchor": "2640", "sl_buffer_amount": "0.5", "sl_buffer_setting": "flat $0.5",
        "params": "[PARAMS] | TF 3m: RR=2 max_SL=$20",
    }
    text = live_history.format_trade_detail(row)
    assert "candle low 2640.00 − buffer 0.50 (flat $0.5) = SL 2639.50" in text
    assert "✓ same as strategy SL" in text
    assert "+ entry offset 0.3 = 2652.30" in text
    assert "TF 3m: RR=2 max_SL=$20" in text


def test_detail_flags_moved_sl_and_shows_error():
    row = {
        "event": "ORDER_FAIL", "direction": "SELL", "reason": "Invalid stops",
        "mt5_message": "10016 Invalid stops", "strategy_sl": "2662", "stop_loss": "2662",
    }
    text = live_history.format_trade_detail(row)
    assert "ERROR / BLOCK" in text and "Invalid stops" in text
    moved = dict(row, event="ORDER", stop_loss="2661.2")
    assert "✗ differs from strategy SL by -0.80" in live_history.format_trade_detail(moved)


# ---------------------------------------------------------------------------
# One line per trade + Excel export
# ---------------------------------------------------------------------------

_BASE = {
    "_journal": "j1", "signal_bar_time": "2026-09-25 10:00", "direction": "BUY",
    "account": "Main", "slot": "Gold", "preset": "gold.json", "pattern": "Hammer",
    "symbol": "XAUUSD", "timeframe": "3m", "signal_high": "2660", "signal_low": "2640",
    "entry_rule": "NEXT_CANDLE_OPEN", "entry_rule_price": "2652", "entry_offset": "0.3",
    "signal_entry": "2652.3", "strategy_entry": "2652.3", "strategy_sl": "2639.5",
    "strategy_tp": "2677.9", "rr_multiple": "2", "sl_anchor_label": "candle low",
    "sl_anchor": "2640", "sl_buffer_amount": "0.5", "sl_buffer_setting": "flat $0.5",
}


def _history_rows():
    return [
        dict(_BASE, event="EXIT", logged_at="2026-09-25T10:40:00", exit_price="2677.3",
             profit="25.2", close_reason="TP", mt5_order_id="778"),
        dict(_BASE, event="ORDER", logged_at="2026-09-25T10:03:01", entry_price="2652.1",
             stop_loss="2639.5", target="2677.3", volume="0.01", mt5_order_id="5512"),
        dict(_BASE, event="SIGNAL", logged_at="2026-09-25T10:03:00"),
        dict(_BASE, event="ORDER_FAIL", logged_at="2026-09-25T09:30:00",
             signal_bar_time="2026-09-25 09:27", reason="10016 Invalid stops"),
        {"_journal": "j1", "event": "ERROR", "logged_at": "2026-09-25T09:00:00",
         "reason": "MT5 disconnected", "account": "Main"},
    ]


def test_calc_texts_spell_out_the_maths():
    r = dict(_BASE, event="ORDER", stop_loss="2639.5")
    assert live_history.entry_calc_text(r) == "next candle open 2652.00 + 0.3 = 2652.30"
    assert live_history.sl_calc_text(r) == "candle low 2640.00 − 0.50 (flat $0.5) = 2639.50"
    assert live_history.tp_calc_text(r) == "risk 12.80 × RR 2 = 25.60 → TP 2677.90"
    assert live_history.sl_check_text(r) == "OK"
    assert live_history.sl_check_text(dict(r, stop_loss="2640.1")) == "MOVED +0.60"


def test_trade_summaries_one_line_per_signal():
    trades = live_history.trade_summaries(_history_rows())
    assert len(trades) == 2  # ERROR without a signal candle only goes to Errors
    closed, failed = trades
    assert closed["Status"] == "Closed"
    assert closed["Entry sent"] == 2652.1 and closed["SL check"] == "OK"
    assert closed["Profit"] == 25.2 and closed["Ticket"] == "5512"
    assert closed["Problem"] == ""
    assert failed["Status"] == "Order failed"
    assert "Invalid stops" in failed["Problem"]


def test_export_workbook_sheets_and_problem_highlight(tmp_path):
    import openpyxl

    path = str(tmp_path / "live.xlsx")
    counts = live_history.export_history_xlsx(path, _history_rows(), filters_text="account=Main")
    assert counts == {"Trades": 2, "Errors": 2, "All_Events": 5}
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["README", "Trades", "Errors", "All_Events"]
    ws = wb["Trades"]
    headers = [c.value for c in ws[1]]
    row2 = {h: c.value for h, c in zip(headers, ws[2])}
    assert row2["SL calc"] == "candle low 2640.00 − 0.50 (flat $0.5) = 2639.50"
    assert row2["Profit"] == 25.2
    fail_fill = ws.cell(row=3, column=1).fill.fgColor.rgb
    assert str(fail_fill).endswith("FDECEA")
    assert wb["README"]["B4"].value == "account=Main"


# ---------------------------------------------------------------------------
# Engine → CSV → History, end to end
# ---------------------------------------------------------------------------

def test_engine_journal_row_carries_breakdown(tmp_path):
    root = str(tmp_path)
    slot_dir = lj.slot_journal_dir(root, "ab12", "Main", "s1", "Gold", "3m")
    cfg = logic.StrategyConfig(
        entry_rule=logic.EntryRule.NEXT_CANDLE_OPEN, entry_offset=0.30,
        buffer_mode=logic.BufferMode.FLAT_AMOUNT, sl_buffer_flat=0.50, enable_risk_limit=False,
    )
    hammer = _c("2026-09-25 10:00", 2650, 2660, 2640, 2655)
    nxt = _c("2026-09-25 10:03", 2652, 2665, 2650, 2660)
    entry = logic.calculate_entry_price(hammer, nxt, cfg)
    sl = logic.calculate_stop_loss(hammer, logic.TradeDirection.BUY, cfg, entry_price=entry)
    sig = _sig(logic.TradeDirection.BUY, hammer, nxt, entry, sl, variant="CLASSIC")

    eng = LiveTradingEngine.__new__(LiveTradingEngine)
    eng.live_config = SimpleNamespace(
        symbol="XAUUSD", timeframe_label="3m", magic=1101, order_mode="market",
        slot_name="Gold", account_name="Main", preset_name="gold.json",
        notify_enabled=False, dry_run=True,
    )
    eng._journal_dir = slot_dir
    eng._active_symbol = "XAUUSD"
    eng.pattern_label = "Hammer"
    eng.pattern_type = "hammer"
    eng.strategy_config = cfg
    eng._telegram = MagicMock()
    eng._record_trade_event("DRY_RUN", sig, volume=0.01, dry_run=True, reason="Dry run")

    rows = live_history.HistoryStore(root).load()
    assert len(rows) == 1
    r = rows[0]
    assert r["preset"] == "gold.json" and r["account"] == "Main"
    assert r["sl_anchor"] == "2640.0" and float(r["sl_buffer_amount"]) == 0.5
    assert float(r["strategy_sl"]) == 2639.5
    assert "RR=" in r["params"]
    text = live_history.format_trade_detail(r)
    assert "candle low 2640.00 − buffer 0.50" in text
