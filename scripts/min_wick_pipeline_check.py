"""
Min-wick filter pipeline check — Hammer with candles + Hammer with candle 35%.

Runs the real backtest on local data and proves, per pattern and timeframe:

  A. Filter OFF = same results as the engine without the filter code at all
     (and the % value is ignored while the switch is off).
  B. Filter ON only removes setups whose wick is below the %, never adds a
     setup and never changes entry / SL / TP of a setup it keeps.
  C. Every signal's entry / SL / TP / pullback limit recomputes by hand, and
     every trade's exit price, size and P&L recompute by hand (all 3 exit models).

It prints each pipeline stage (candles → setups → min-wick → risk → fills →
trades → results) and writes an Excel workbook with every setup and trade.

Usage (from repo root):
    python scripts/min_wick_pipeline_check.py
    python scripts/min_wick_pipeline_check.py --timeframes 1hour,5min --year 2024
    python scripts/min_wick_pipeline_check.py --buy-pct 40 --sell-pct 35 --sides both
    python scripts/min_wick_pipeline_check.py --wick-shape off        # body-only mode
    python scripts/min_wick_pipeline_check.py --preset presets/my_setup.json

Exit code 0 = every check passed, 1 = at least one FAIL.
"""

from __future__ import annotations

import argparse
import contextlib
import copy
import hashlib
import io
import json
import os
import sys
from collections import Counter
from datetime import datetime
from typing import Dict, List, Optional, Tuple

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import logic  # noqa: E402
import hammer_context_logic as hc  # noqa: E402
from strategy import hammer_context_core as core  # noqa: E402
from strategy import hammer_with_candles_35_logic as hwc35  # noqa: E402
from strategy import hammer_with_candles_logic as hwc  # noqa: E402
from backtest import (  # noqa: E402
    ALL_EXIT_MODELS,
    BacktestConfig,
    ExitModel,
    PositionSizingMode,
    TradeOutcome,
    df_to_candles,
    load_candles_df,
    market_preferred_for_config,
    run_full_backtest,
)

TOL = 1e-6
BUY = logic.TradeDirection.BUY
PRODUCTS = [
    ("Hammer with candles", hwc.PATTERN_TYPE, hwc),
    ("Hammer with candle 35%", hwc35.PATTERN_TYPE, hwc35),
]


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

def base_strategy(args, label: str) -> hc.HammerContextConfig:
    if args.preset:
        from API.runtime import build_hammer_context_strategy

        with open(args.preset, "r", encoding="utf-8") as f:
            data = dict(json.load(f))
        data["pattern"] = label
        return build_hammer_context_strategy(data)
    pull = hwc35.DEFAULT_PULLBACK_PCT if hwc35.is_pattern_label(label) else 0.0
    wick_on = args.wick_shape == "on"
    return hc.HammerContextConfig(
        entry_pullback_pct=pull, buy_require_wick=wick_on, sell_require_wick=wick_on,
    )


def with_filter(base, *, buy_on: bool, sell_on: bool, buy_pct: float, sell_pct: float):
    cfg = copy.deepcopy(base)
    cfg.buy_min_lower_wick_enabled = buy_on
    cfg.sell_min_upper_wick_enabled = sell_on
    cfg.buy_min_lower_wick_pct = float(buy_pct)
    cfg.sell_min_upper_wick_pct = float(sell_pct)
    return cfg


def backtest_config(args, ptype: str, strat, tf: str) -> BacktestConfig:
    return BacktestConfig(
        strategy_config=strat,
        pattern_type=ptype,
        symbol=args.symbol,
        data_root=args.data_root,
        start_date=datetime(args.year, 1, 1),
        end_date=datetime(args.year + 1, 1, 1),
        timeframes_to_test=[tf],
        max_forward_candles=200,
        position_sizing_mode=PositionSizingMode.FIXED_RISK_USD,
        fixed_risk_usd=args.risk_usd,
        starting_capital=10000.0,
        allow_overlapping_trades=False,
        sessions_enabled={"Asian": True, "London": True, "US": True},
    )


def logic_label_for(cfg: BacktestConfig, tf: str) -> str:
    label = cfg.timeframe_folder_to_logic_label.get(tf)
    if label:
        return label
    try:
        from candle_resample import logic_label_for_folder

        return logic_label_for_folder(tf)
    except Exception:
        return tf


# ---------------------------------------------------------------------------
# Independent maths (written from the rules, not by calling the engine)
# ---------------------------------------------------------------------------

def wick_pcts(c: logic.Candle) -> Tuple[float, float]:
    rng = c.high - c.low
    if rng <= 1e-9:
        return 0.0, 0.0
    return (min(c.open, c.close) - c.low) / rng * 100.0, (c.high - max(c.open, c.close)) / rng * 100.0


def side_params(strat, is_buy: bool):
    if is_buy:
        return (strat.entry_rule, strat.entry_offset, strat.sl_mode, strat.sl_fixed_distance,
                strat.buffer_mode, strat.sl_buffer_pct, strat.sl_buffer_flat)
    return (strat.inverted_entry_rule, strat.inverted_entry_offset, strat.inverted_sl_mode,
            strat.inverted_sl_fixed_distance, strat.inverted_buffer_mode,
            strat.inverted_sl_buffer_pct, strat.inverted_sl_buffer_flat)


def expected_signal(sig, candles, idx: int, strat, label: str) -> Dict[str, float]:
    is_buy = sig.direction == BUY
    h, n = candles[idx], candles[idx + 1]
    rule, offset, sl_mode, fixed, buf_mode, buf_pct, buf_flat = side_params(strat, is_buy)
    rule = logic.coerce_entry_rule(rule)
    base = {
        logic.EntryRule.NEXT_CANDLE_OPEN: n.open,
        logic.EntryRule.NEXT_CANDLE_CLOSE: n.close,
        logic.EntryRule.HAMMER_CLOSE: h.close,
        logic.EntryRule.HAMMER_HIGH: h.high,
        logic.EntryRule.HAMMER_LOW: h.low,
    }[rule] + float(offset or 0.0)

    if logic.coerce_stop_loss_mode(sl_mode) == logic.StopLossMode.FIXED_FROM_ENTRY:
        dist = max(0.0, float(fixed or 0.0))
        sl = base - dist if is_buy else base + dist
    else:
        anchor = h.low if is_buy else h.high
        mode = logic.coerce_buffer_mode(buf_mode)
        buf = {
            logic.BufferMode.PERCENT_OF_RANGE: (h.high - h.low) * float(buf_pct or 0.0) / 100.0,
            logic.BufferMode.PERCENT_OF_PRICE: anchor * float(buf_pct or 0.0) / 100.0,
            logic.BufferMode.FLAT_AMOUNT: float(buf_flat or 0.0),
            logic.BufferMode.NONE: 0.0,
        }[mode]
        sl = anchor - buf if is_buy else anchor + buf

    sig_risk = base - sl if is_buy else sl - base
    tf_set = logic.resolve_timeframe_setting(strat.timeframe_settings, label)
    target = base + sig_risk * tf_set.rr_multiple if is_buy else base - sig_risk * tf_set.rr_multiple
    pull = float(strat.entry_pullback_pct or 0.0)
    entry = base
    if pull > 0:
        shift = abs(base - sl) * min(100.0, pull) / 100.0
        entry = base - shift if is_buy else base + shift
    fill_risk = entry - sl if is_buy else sl - entry
    ignored = (strat.reject_zero_or_negative_risk and sig_risk <= 0) or (
        not (strat.reject_zero_or_negative_risk and sig_risk <= 0)
        and strat.enable_risk_limit and sig_risk > tf_set.max_sl_usd
    )
    return {
        "base_entry": base, "entry": entry, "sl": sl, "target": target,
        "risk": fill_risk if fill_risk > 0 else sig_risk, "ignored": bool(ignored),
        "rr": tf_set.rr_multiple,
    }


def check_trade(t, cfg: BacktestConfig) -> List[str]:
    """Recompute exit/size/P&L for one ledger row under every exit model."""
    sig = t.signal
    is_buy = sig.direction == BUY
    errs: List[str] = []
    entry, sl, target = float(sig.entry_price), float(sig.stop_loss), float(sig.target)
    risk = entry - sl if is_buy else sl - entry
    if abs(risk - float(sig.risk)) > TOL:
        errs.append(f"risk {sig.risk:.5f} != |entry-SL| {risk:.5f}")
    if getattr(sig, "await_limit_fill", False) and sig.signal_entry_price is not None:
        base = float(sig.signal_entry_price)
        limit = base - abs(base - sl) * float(sig.entry_pullback_pct) / 100.0 if is_buy \
            else base + abs(base - sl) * float(sig.entry_pullback_pct) / 100.0
        # Gap fills may only improve the price (never worse than the limit, never past SL)
        if (is_buy and not (sl < entry <= limit + TOL)) or (not is_buy and not (limit - TOL <= entry < sl)):
            errs.append(f"fill {entry:.5f} outside limit {limit:.5f} / SL {sl:.5f}")

    for m in ALL_EXIT_MODELS:
        outcome, px, _ts, _bars = t.outcomes[m]
        pnl = float(t.pnl.get(m, 0.0))
        size = float(t.position_size.get(m, 0.0))
        tag = m.value
        if outcome in (TradeOutcome.WIN, TradeOutcome.LOSS):
            if outcome == TradeOutcome.WIN and abs(float(px) - target) > TOL:
                errs.append(f"{tag}: WIN exit {px} != TP {target:.5f}")
            if outcome == TradeOutcome.LOSS and ((is_buy and px > sl + TOL) or (not is_buy and px < sl - TOL)):
                errs.append(f"{tag}: LOSS exit {px} better than SL {sl:.5f}")
            exp_size = cfg.fixed_risk_usd / risk if risk > 0 else 0.0
            if abs(size - exp_size) > 1e-6 * max(1.0, exp_size):
                errs.append(f"{tag}: size {size:.6f} != risk$/risk {exp_size:.6f}")
            slip = cfg.slippage_usd
            e_adj = entry + slip if is_buy else entry - slip
            x_adj = float(px) - slip if is_buy else float(px) + slip
            exp_pnl = ((x_adj - e_adj) if is_buy else (e_adj - x_adj)) * exp_size - cfg.commission_per_trade
            if abs(pnl - exp_pnl) > 1e-4:
                errs.append(f"{tag}: pnl {pnl:.4f} != recomputed {exp_pnl:.4f}")
        elif abs(pnl) > TOL:
            errs.append(f"{tag}: {outcome.name} must book 0 P&L, got {pnl}")
    return errs


# ---------------------------------------------------------------------------
# Running
# ---------------------------------------------------------------------------

def signal_key(sig) -> Tuple[str, str]:
    return str(sig.hammer_candle.timestamp), sig.direction.value


def signal_fields(sig) -> tuple:
    return (round(sig.entry_price, 6), round(sig.stop_loss, 6), round(sig.target, 6),
            round(sig.risk, 6), bool(sig.ignored), sig.ignore_reason or "",
            bool(sig.await_limit_fill), sig.entry_pullback_pct)


def fingerprint(ledger) -> str:
    h = hashlib.sha1()
    for t in ledger:
        for m in ALL_EXIT_MODELS:
            o, px, ts, bars = t.outcomes[m]
            h.update(
                f"{m.value}|{t.entry_time}|{t.signal.direction.value}|{t.signal.entry_price:.5f}|"
                f"{t.signal.stop_loss:.5f}|{t.signal.target:.5f}|{o.name}|{px}|{ts}|{bars}|"
                f"{t.pnl.get(m, 0.0):.4f}\n".encode()
            )
    return h.hexdigest()[:12]


def run_backtest(args, ptype, strat, tf):
    cfg = backtest_config(args, ptype, copy.deepcopy(strat), tf)
    with contextlib.redirect_stdout(io.StringIO()):
        ledger, ignored = run_full_backtest(cfg)
    return cfg, ledger, ignored


@contextlib.contextmanager
def filter_code_removed():
    """Reference engine: the min-wick check does not exist at all."""
    saved = core.min_wick_ok
    core.min_wick_ok = lambda *_a, **_k: (True, "")
    try:
        yield
    finally:
        core.min_wick_ok = saved


def results_line(ledger, model: ExitModel) -> Dict[str, object]:
    cnt = Counter(t.outcomes[model][0].name for t in ledger)
    net = sum(float(t.pnl.get(model, 0.0)) for t in ledger)
    closed = cnt.get("WIN", 0) + cnt.get("LOSS", 0)
    return {
        "trades": closed, "wins": cnt.get("WIN", 0), "losses": cnt.get("LOSS", 0),
        "open": cnt.get("STILL_OPEN", 0), "overlap_skipped": cnt.get("SKIPPED_OVERLAP", 0),
        "win_pct": (cnt.get("WIN", 0) / closed * 100.0) if closed else 0.0, "net": net,
    }


class Report:
    def __init__(self):
        self.checks: List[Dict[str, str]] = []
        self.summary: List[Dict[str, object]] = []
        self.setups: List[Dict[str, object]] = []
        self.trades: List[Dict[str, object]] = []

    def check(self, where: str, name: str, ok: bool, detail: str = "") -> None:
        self.checks.append({"where": where, "check": name, "result": "PASS" if ok else "FAIL", "detail": detail})
        print(f"    [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))

    @property
    def failed(self) -> int:
        return sum(1 for c in self.checks if c["result"] == "FAIL")


def run_case(args, rep: Report, label: str, ptype: str, module, tf: str) -> None:
    base = base_strategy(args, label)
    buy_on = args.sides in ("both", "buy")
    sell_on = args.sides in ("both", "sell")
    buy_pct = args.buy_pct if args.buy_pct is not None else (
        base.buy_min_lower_wick_pct if base.buy_min_lower_wick_enabled else 35.0)
    sell_pct = args.sell_pct if args.sell_pct is not None else (
        base.sell_min_upper_wick_pct if base.sell_min_upper_wick_enabled else 35.0)
    cfg_off = with_filter(base, buy_on=False, sell_on=False, buy_pct=35.0, sell_pct=35.0)
    cfg_off_odd = with_filter(base, buy_on=False, sell_on=False, buy_pct=99.0, sell_pct=99.0)
    cfg_on = with_filter(base, buy_on=buy_on, sell_on=sell_on, buy_pct=buy_pct, sell_pct=sell_pct)
    where = f"{label} | {tf}"
    on_desc = " + ".join(
        p for p in (
            f"BUY lower ≥ {buy_pct:g}%" if buy_on else "",
            f"SELL upper ≥ {sell_pct:g}%" if sell_on else "",
        ) if p
    )

    probe_cfg = backtest_config(args, ptype, cfg_off, tf)
    label_tf = logic_label_for(probe_cfg, tf)
    df = load_candles_df(args.data_root, args.symbol, tf, probe_cfg.start_date, probe_cfg.end_date,
                         preferred=market_preferred_for_config(probe_cfg))
    candles = df_to_candles(df)
    print(f"\n=== {label} | {tf} ({label_tf}) | {args.year} | filter ON = {on_desc} ===")
    if len(candles) < 3:
        rep.check(where, "Candle data present", False, f"only {len(candles)} candles")
        return
    ts_index = {str(c.timestamp): i for i, c in enumerate(candles)}
    prepared_on = module.prepare_config(cfg_on)

    # ---- Signal stage (same call the backtest makes) ----
    sig_off = module.run_strategy(candles, label_tf, cfg_off)
    sig_odd = module.run_strategy(candles, label_tf, cfg_off_odd)
    sig_on = module.run_strategy(candles, label_tf, cfg_on)
    with filter_code_removed():
        sig_ref = module.run_strategy(candles, label_tf, cfg_on)

    off_map = {signal_key(s): s for s in sig_off}
    on_map = {signal_key(s): s for s in sig_on}
    dir_cnt = Counter(s.direction.value for s in sig_off)
    removed = [k for k in off_map if k not in on_map]
    removed_dir = Counter(k[1] for k in removed)
    added = [k for k in on_map if k not in off_map]

    print(f"  1. Candles loaded ............. {len(candles):,}")
    print(f"  2. Setups found (filter OFF) .. {len(sig_off):,}  (BUY {dir_cnt.get('BUY', 0)} / SELL {dir_cnt.get('SELL', 0)})")
    print(f"  3. Min-wick filter ON ......... removed {len(removed):,} "
          f"(BUY {removed_dir.get('BUY', 0)} / SELL {removed_dir.get('SELL', 0)}) → kept {len(sig_on):,}")
    risk_off = sum(1 for s in sig_off if s.ignored)
    risk_on = sum(1 for s in sig_on if s.ignored)
    print(f"  4. Risk filter (max SL) ....... skipped OFF {risk_off} / ON {risk_on}")

    # Per-setup expectations
    wick_mismatch, calc_errs, kept_changed = [], [], []
    for key, s in off_map.items():
        idx = ts_index.get(key[0])
        is_buy = s.direction == BUY
        lower, upper = wick_pcts(s.hammer_candle)
        wick = lower if is_buy else upper
        side_on = buy_on if is_buy else sell_on
        need = buy_pct if is_buy else sell_pct
        should_keep = (not side_on) or wick >= need
        kept = key in on_map
        if should_keep != kept:
            wick_mismatch.append(f"{key[0]} {key[1]} wick {wick:.3f}% need {need:g}%")
        if kept and signal_fields(on_map[key]) != signal_fields(s):
            kept_changed.append(f"{key[0]} {key[1]}")
        exp = expected_signal(s, candles, idx, prepared_on, label_tf) if idx is not None else None
        bad = []
        if exp is None:
            bad.append("signal bar not in candles")
        else:
            for name, got in (("entry", s.entry_price), ("sl", s.stop_loss), ("target", s.target), ("risk", s.risk)):
                if abs(float(got) - exp[name]) > TOL:
                    bad.append(f"{name} {got:.5f} != {exp[name]:.5f}")
            if bool(s.ignored) != exp["ignored"]:
                bad.append(f"ignored {s.ignored} != {exp['ignored']}")
        if bad:
            calc_errs.append(f"{key[0]} {key[1]}: " + "; ".join(bad))
        rep.setups.append({
            "pattern": label, "timeframe": tf, "signal_time": key[0], "direction": key[1],
            "open": s.hammer_candle.open, "high": s.hammer_candle.high,
            "low": s.hammer_candle.low, "close": s.hammer_candle.close,
            "range": s.hammer_candle.high - s.hammer_candle.low,
            "lower_wick_pct": round(lower, 3), "upper_wick_pct": round(upper, 3),
            "checked_wick": "lower" if is_buy else "upper", "min_pct": need if side_on else None,
            "filter_ON": "KEPT" if kept else "REMOVED",
            "signal_entry": exp["base_entry"] if exp else None,
            "entry": s.entry_price, "sl": s.stop_loss, "tp": s.target, "risk": s.risk,
            "rr": exp["rr"] if exp else None,
            "risk_skipped": s.ignore_reason or "", "calc_check": "OK" if not bad else "; ".join(bad),
        })

    rep.check(where, "Setups: switch OFF ignores the % (35 vs 99 identical)",
              [signal_fields(s) for s in sig_off] == [signal_fields(s) for s in sig_odd]
              and [signal_key(s) for s in sig_off] == [signal_key(s) for s in sig_odd])
    rep.check(where, "Setups: OFF identical to engine without filter code",
              [(signal_key(s), signal_fields(s)) for s in sig_off]
              == [(signal_key(s), signal_fields(s)) for s in sig_ref])
    rep.check(where, "Setups: ON never adds a setup", not added, f"{len(added)} added" if added else "")
    rep.check(where, "Setups: ON removes exactly the setups below the wick %", not wick_mismatch,
              "; ".join(wick_mismatch[:5]) if wick_mismatch else (
                  f"{len(removed)} removed (all below the %), {len(sig_on)} kept (all at/above)"
                  if removed else "0 removed — every setup already meets the %"))
    rep.check(where, "Setups: kept setups have unchanged entry/SL/TP/risk", not kept_changed,
              ", ".join(kept_changed[:5]))
    rep.check(where, "Setups: entry / SL / TP / risk recompute by hand", not calc_errs,
              "; ".join(calc_errs[:3]) if calc_errs else f"{len(off_map)} setups checked")

    # ---- Backtest stage ----
    cfg_bt_off, led_off, ign_off = run_backtest(args, ptype, cfg_off, tf)
    _c, led_odd, _i = run_backtest(args, ptype, cfg_off_odd, tf)
    cfg_bt_on, led_on, ign_on = run_backtest(args, ptype, cfg_on, tf)
    with filter_code_removed():
        _c, led_ref, _i = run_backtest(args, ptype, cfg_on, tf)

    nf_off = sum(1 for s in ign_off if "not filled" in (s.ignore_reason or ""))
    nf_on = sum(1 for s in ign_on if "not filled" in (s.ignore_reason or ""))
    if getattr(prepared_on, "entry_pullback_pct", 0.0):
        print(f"  5. Pullback limit {prepared_on.entry_pullback_pct:g}% not filled  OFF {nf_off} / ON {nf_on}")
    else:
        print("  5. Entry ...................... market (no pullback)")
    m = ExitModel.WORST_CASE
    r_off, r_on = results_line(led_off, m), results_line(led_on, m)
    print(f"  6. Trades (worst case) ........ OFF {r_off['trades']} (overlap skipped {r_off['overlap_skipped']}, open {r_off['open']})"
          f" / ON {r_on['trades']} (overlap skipped {r_on['overlap_skipped']}, open {r_on['open']})")
    print(f"  7. Results (worst case) ....... OFF W{r_off['wins']}/L{r_off['losses']} win {r_off['win_pct']:.1f}% net ${r_off['net']:,.2f}"
          f"  |  ON W{r_on['wins']}/L{r_on['losses']} win {r_on['win_pct']:.1f}% net ${r_on['net']:,.2f}")

    fp_off, fp_odd, fp_ref, fp_on = fingerprint(led_off), fingerprint(led_odd), fingerprint(led_ref), fingerprint(led_on)
    rep.check(where, "Backtest: switch OFF ignores the % (35 vs 99 identical)", fp_off == fp_odd, f"sha {fp_off}")
    rep.check(where, "Backtest: OFF identical to engine without filter code", fp_off == fp_ref, f"sha {fp_off}")
    if removed:
        rep.check(where, "Backtest: ON differs from OFF (filter removed setups)", fp_on != fp_off,
                  f"sha OFF {fp_off} → ON {fp_on}")
    else:
        rep.check(where, "Backtest: ON equals OFF (no setup was below the %)", fp_on == fp_off, f"sha {fp_off}")

    for mode, ledger, cfg_bt in (("OFF", led_off, cfg_bt_off), ("ON", led_on, cfg_bt_on)):
        errs_total = []
        for t in ledger:
            errs = check_trade(t, cfg_bt)
            if errs:
                errs_total.append(f"{t.entry_time} {t.signal.direction.value}: " + "; ".join(errs))
            o, px, ets, bars = t.outcomes[m]
            rep.trades.append({
                "pattern": label, "timeframe": tf, "filter": mode,
                "signal_time": str(t.signal.hammer_candle.timestamp), "entry_time": str(t.entry_time),
                "direction": t.signal.direction.value,
                "signal_entry": t.signal.signal_entry_price if t.signal.signal_entry_price is not None else t.signal.entry_price,
                "entry": t.signal.entry_price, "sl": t.signal.stop_loss, "tp": t.signal.target,
                "risk": t.signal.risk, "outcome_worst": o.name, "exit_price": px, "exit_time": str(ets or ""),
                "bars": bars, "size": t.position_size.get(m, 0.0), "risk_usd": t.risk_usd.get(m, 0.0),
                "pnl_worst": t.pnl.get(m, 0.0), "pnl_best": t.pnl.get(ExitModel.BEST_CASE, 0.0),
                "pnl_candle_bias": t.pnl.get(ExitModel.CANDLE_BIAS, 0.0),
                "math_check": "OK" if not errs else "; ".join(errs),
            })
        rep.check(where, f"Trades {mode}: exit / size / P&L recompute by hand (3 exit models)", not errs_total,
                  "; ".join(errs_total[:3]) if errs_total else f"{len(ledger)} ledger rows checked")

    for mode, r in (("OFF", r_off), ("ON", r_on)):
        rep.summary.append({
            "pattern": label, "timeframe": tf, "year": args.year, "filter": mode,
            "filter_rule": on_desc if mode == "ON" else "off",
            "candles": len(candles),
            "setups": len(sig_off) if mode == "OFF" else len(sig_on),
            "removed_by_min_wick": 0 if mode == "OFF" else len(removed),
            "risk_skipped": risk_off if mode == "OFF" else risk_on,
            "limit_not_filled": nf_off if mode == "OFF" else nf_on,
            **{k: (round(v, 2) if isinstance(v, float) else v) for k, v in r.items()},
            "fingerprint": fp_off if mode == "OFF" else fp_on,
        })


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def write_excel(path: str, rep: Report, args) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="34A853")
    bad_fill = PatternFill("solid", fgColor="FDECEA")
    ok_fill = PatternFill("solid", fgColor="E6F4EA")

    def sheet(name: str, rows: List[Dict[str, object]], flag_col: Optional[str] = None, ok_value: str = "OK"):
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["(no rows)"])
            return
        headers = list(rows[0].keys())
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append([r.get(h) for h in headers])
            if flag_col:
                val = str(r.get(flag_col, ""))
                if val and val not in (ok_value, "PASS", "KEPT"):
                    for c in ws[ws.max_row]:
                        c.fill = bad_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in list(col)[:200])
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))

    readme = wb.active
    readme.title = "README"
    lines = [
        "Min-wick filter pipeline check",
        f"Generated {datetime.now():%Y-%m-%d %H:%M}  |  data {args.data_root}  |  year {args.year}"
        f"  |  timeframes {args.timeframes}  |  preset {args.preset or '(defaults)'}"
        f"  |  wick shape {'from preset' if args.preset else args.wick_shape}",
        f"Result: {len(rep.checks) - rep.failed} PASS / {rep.failed} FAIL",
        "",
        "Checks — every PASS/FAIL line (FAIL rows are red).",
        "Summary — pipeline counts and results per pattern / timeframe, filter OFF vs ON (worst-case model).",
        "Setups — every setup found with filter OFF: wick %s, KEPT/REMOVED with filter ON, hand-checked entry/SL/TP.",
        "Trades — every trade OFF and ON: entry/SL/TP/exit/size/P&L with a hand recompute (math_check).",
        "Fingerprint = hash of every trade under all 3 exit models: equal fingerprints = identical backtests.",
    ]
    for line in lines:
        readme.append([line])
    readme["A1"].font = Font(bold=True, size=14)
    readme["A3"].fill = ok_fill if rep.failed == 0 else bad_fill

    sheet("Checks", rep.checks, flag_col="result", ok_value="PASS")
    sheet("Summary", rep.summary)
    sheet("Setups", rep.setups, flag_col="calc_check")
    sheet("Trades", rep.trades, flag_col="math_check")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data-root", default=os.path.join(ROOT, "data"))
    ap.add_argument("--symbol", default="XAUUSD")
    ap.add_argument("--timeframes", default="1hour,5min", help="comma list of data folders")
    ap.add_argument("--year", type=int, default=2024)
    ap.add_argument("--sides", choices=("both", "buy", "sell"), default="both",
                    help="which side(s) get the filter in the ON run")
    ap.add_argument("--buy-pct", type=float, default=None, help="min lower wick %% for BUY (default 35)")
    ap.add_argument("--sell-pct", type=float, default=None, help="min upper wick %% for SELL (default 35)")
    ap.add_argument("--risk-usd", type=float, default=100.0, help="fixed $ risk per trade")
    ap.add_argument("--preset", default="", help="preset JSON to use for all other settings")
    ap.add_argument("--wick-shape", choices=("on", "off"), default="on",
                    help="Require wick shape? (ignored with --preset, which has its own)")
    ap.add_argument("--patterns", default="both", choices=("both", "hwc", "hwc35"))
    ap.add_argument("--excel", default="", help="output .xlsx (default output/checks/min_wick_check_<time>.xlsx)")
    args = ap.parse_args()

    rep = Report()
    products = [p for p in PRODUCTS if args.patterns == "both"
                or (args.patterns == "hwc" and p[2] is hwc) or (args.patterns == "hwc35" and p[2] is hwc35)]
    for tf in [t.strip() for t in args.timeframes.split(",") if t.strip()]:
        for label, ptype, module in products:
            try:
                run_case(args, rep, label, ptype, module, tf)
            except Exception as e:
                rep.check(f"{label} | {tf}", "Run completed", False, f"{type(e).__name__}: {e}")

    out = args.excel or os.path.join(ROOT, "output", "checks", f"min_wick_check_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    write_excel(out, rep, args)
    print(f"\n{'=' * 70}\nRESULT: {len(rep.checks) - rep.failed} PASS / {rep.failed} FAIL")
    print(f"Excel: {out}")
    return 1 if rep.failed else 0


if __name__ == "__main__":
    sys.exit(main())
